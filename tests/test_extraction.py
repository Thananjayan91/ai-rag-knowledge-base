import pytest
from docx import Document as DocxDocument

from backend.extraction import extract_document, extract_docx, extract_txt


def test_extract_txt_reads_file_content(tmp_path):
    f = tmp_path / "sample.txt"
    f.write_text("hello world", encoding="utf-8")
    pages = extract_txt(str(f))
    assert len(pages) == 1
    assert pages[0].page is None
    assert "hello world" in pages[0].text


def test_extract_txt_empty_file_returns_no_pages(tmp_path):
    f = tmp_path / "empty.txt"
    f.write_text("", encoding="utf-8")
    assert extract_txt(str(f)) == []


def test_extract_docx_reads_paragraphs_in_order(tmp_path):
    f = tmp_path / "sample.docx"
    doc = DocxDocument()
    doc.add_paragraph("First paragraph")
    doc.add_paragraph("Second paragraph")
    doc.save(str(f))

    pages = extract_docx(str(f))
    assert len(pages) == 1
    text = pages[0].text
    assert text.index("First paragraph") < text.index("Second paragraph")


def test_extract_document_dispatches_by_extension(tmp_path):
    f = tmp_path / "sample.txt"
    f.write_text("routed correctly", encoding="utf-8")
    pages = extract_document(str(f), "sample.txt")
    assert "routed correctly" in pages[0].text


def test_extract_document_rejects_unsupported_type(tmp_path):
    f = tmp_path / "sample.xyz"
    f.write_text("data", encoding="utf-8")
    with pytest.raises(ValueError):
        extract_document(str(f), "sample.xyz")
